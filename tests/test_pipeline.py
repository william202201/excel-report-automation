import csv
from datetime import datetime
from decimal import Decimal
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from excel_report.cli import main
from excel_report.demo import generate
from excel_report.pipeline import consolidate, number
from excel_report.report import write_report


class PipelineTest(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "input"
        self.source.mkdir()

    def fixture(self, rows, name="one.xlsx", header=None):
        wb = Workbook()
        ws = wb.active
        ws.append(["模拟数据"])
        ws.append(header or ["订单号", "日期", "区域", "商品名称", "销售数量", "销售单价", "金额"])
        for row in rows:
            ws.append(row)
        wb.save(self.source / name)

    def row(self, key="A1", **changes):
        result = [key, "2025/01/02", "East", "键盘", 2, "￥49.50", 99]
        for index, value in changes.items():
            result[int(index)] = value
        return result

    def test_normalization_and_exact_money(self):
        self.fixture([self.row(), self.row("A2", **{"1": datetime(2025, 2, 3), "4": 3, "5": "0.10", "6": "0.30"})])
        result = consolidate(self.source)
        self.assertEqual(result.metrics["sales_total_cny"], "99.30")
        self.assertEqual(result.rows[0]["region"], "华东")
        self.assertEqual(result.rows[0]["row"], 3)
        self.assertTrue(result.metrics["reconciled"])

    def test_exact_duplicate_across_files(self):
        self.fixture([self.row()])
        self.fixture([self.row()], "two.xlsx")
        result = consolidate(self.source)
        self.assertEqual((len(result.rows), result.duplicates, result.quarantined), (1, 1, 0))

    def test_conflicting_id_quarantines_both(self):
        self.fixture([self.row(), self.row(**{"3": "鼠标"})])
        result = consolidate(self.source)
        self.assertEqual((len(result.rows), result.quarantined), (0, 2))
        self.assertTrue(all(r["reason"] == "conflicting:order_id" for r in result.issues))

    def test_bad_values_are_not_silently_fixed(self):
        cases = [(1, "2025-02-30"), (2, "unknown"), (3, None), (4, -2), (4, 1.5), (5, "NaN"), (5, "Infinity"), (5, "1,23"), (6, 0)]
        self.fixture([self.row(f"A{i}", **{str(field): value}) for i, (field, value) in enumerate(cases)])
        result = consolidate(self.source)
        self.assertEqual(result.quarantined, len(cases))
        self.assertEqual(result.metrics["sales_total_cny"], "0.00")

    def test_amount_can_be_derived_when_column_absent(self):
        self.fixture([self.row()[:6]], header=["订单号", "日期", "区域", "商品名称", "销售数量", "销售单价"])
        self.assertEqual(consolidate(self.source).rows[0]["amount"], Decimal("99.00"))

    def test_unknown_schema_fails_closed(self):
        self.fixture([[1, 2]], header=["bad", "columns"])
        with self.assertRaisesRegex(ValueError, "No valid header"):
            consolidate(self.source)

    def test_ambiguous_headers_fail(self):
        self.fixture([self.row() + ["A1"]], header=["订单号", "日期", "区域", "商品名称", "销售数量", "销售单价", "金额", "Order ID"])
        with self.assertRaisesRegex(ValueError, "Ambiguous"):
            consolidate(self.source)

    def test_reordered_columns(self):
        self.fixture([list(reversed(self.row()))], header=list(reversed(["订单号", "日期", "区域", "商品名称", "销售数量", "销售单价", "金额"])))
        self.assertEqual(consolidate(self.source).metrics["sales_total_cny"], "99.00")

    def test_saved_report_charts_and_reconciliation(self):
        self.fixture([self.row(), self.row("A2", **{"4": -1})])
        output = self.root / "output"
        write_report(consolidate(self.source), output)
        wb = load_workbook(output / "report.xlsx")
        self.addCleanup(wb.close)
        self.assertEqual(len(wb.sheetnames), 7)
        self.assertEqual(len(wb["总览"]._charts), 2)
        self.assertEqual(wb["月度汇总"]["B2"].value, 99)
        self.assertEqual(wb["地区产品透视汇总"]["C3"].value, 99)
        with (output / "issues.csv").open(encoding="utf-8-sig", newline="") as handle:
            self.assertEqual(len(list(csv.DictReader(handle))), 1)
        self.assertTrue(json.loads((output / "summary.json").read_text())["reconciled"])

    def test_all_bad_rows_still_produce_empty_report(self):
        self.fixture([self.row(**{"4": -1})])
        write_report(consolidate(self.source), self.root / "empty")
        self.assertIn("0.00", (self.root / "empty" / "summary.json").read_text())

    def test_formula_injection_is_stored_as_text(self):
        self.fixture([self.row(**{"3": '=HYPERLINK("https://example.invalid")'})])
        write_report(consolidate(self.source), self.root / "safe")
        wb = load_workbook(self.root / "safe" / "report.xlsx", data_only=False)
        self.addCleanup(wb.close)
        self.assertEqual(wb["清洗明细"]["D2"].data_type, "s")

    def test_demo_has_12_formats_and_known_accounting(self):
        self.source.rmdir()
        generate(self.source, 60)
        result = consolidate(self.source)
        self.assertEqual(result.metrics["input_files"], 12)
        self.assertEqual((result.raw_rows, len(result.rows), result.quarantined, result.duplicates), (744, 636, 96, 12))

    def test_missing_input_returns_nonzero(self):
        self.assertEqual(main(["run", "--input", str(self.root / "missing")]), 2)

    def test_numeric_grouping_and_nonfinite(self):
        self.assertEqual(number("￥1,234.50"), Decimal("1234.50"))
        for value in (True, None, "NaN", "-Infinity", "1,23.45"):
            with self.subTest(value=value), self.assertRaises(ValueError):
                number(value)


if __name__ == "__main__":
    unittest.main()
