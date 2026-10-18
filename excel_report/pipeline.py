"""Read every worksheet, validate row-level values, reconcile every input row."""
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import json
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook

FIELDS = ["order_id", "date", "region", "product", "quantity", "unit_price", "amount"]
CENT = Decimal("0.01")


def norm(value) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(value or ""))).casefold()


def load_schema(path: Path | None = None) -> dict:
    return json.loads((path or Path(__file__).with_name("schema.json")).read_text(encoding="utf-8"))


def number(value) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise ValueError("invalid number")
    text = unicodedata.normalize("NFKC", str(value)).strip()
    # Commas must be valid thousands separators; do not guess decimal locale.
    if "," in text and not re.fullmatch(r"[￥¥]?\s*[+-]?\d{1,3}(,\d{3})+(\.\d+)?", text):
        raise ValueError("invalid thousands grouping")
    try:
        result = Decimal(text.removeprefix("￥").removeprefix("¥").replace(",", "").strip())
    except InvalidOperation as exc:
        raise ValueError("invalid number") from exc
    if not result.is_finite():
        raise ValueError("non-finite number")
    return result


def parse_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y年%m月%d日"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    raise ValueError("invalid date")


@dataclass
class Result:
    rows: list[dict]
    issues: list[dict]
    sources: list[dict]
    raw_rows: int
    duplicates: int
    quarantined: int

    @property
    def metrics(self):
        return {"input_files": len({s["file"] for s in self.sources}),
                "raw_rows": self.raw_rows, "accepted_rows": len(self.rows),
                "quarantined_rows": self.quarantined, "duplicate_rows": self.duplicates,
                "sales_total_cny": str(sum((r["amount"] for r in self.rows), Decimal(0)).quantize(CENT)),
                "reconciled": self.raw_rows == len(self.rows) + self.duplicates + self.quarantined}


def consolidate(folder: Path, schema: dict | None = None) -> Result:
    schema = schema or load_schema()
    aliases = {norm(alias): field for field, values in schema["fields"].items() for alias in [field, *values]}
    files = sorted(p for p in folder.glob("*.xlsx") if not p.name.startswith("~$"))
    if not files:
        raise ValueError(f"No .xlsx inputs found in {folder}")
    candidates, issues, sources = [], [], []
    raw = quarantined = duplicates = 0
    for file in files:
        wb = load_workbook(file, read_only=True, data_only=False)
        matched = 0
        try:
            for ws in wb:
                header = mapping = None
                for index, values in enumerate(ws.iter_rows(values_only=True), 1):
                    if index > 20:
                        break
                    found = [(i, aliases[norm(v)]) for i, v in enumerate(values) if norm(v) in aliases]
                    if set(schema["required"]).issubset({field for _, field in found}):
                        if len({f for _, f in found}) != len(found):
                            raise ValueError(f"Ambiguous duplicate columns: {file.name}/{ws.title}")
                        header, mapping = index, dict(found)
                        break
                if header is None:
                    continue
                matched += 1
                count = 0
                for index, values in enumerate(ws.iter_rows(min_row=header + 1, values_only=True), header + 1):
                    if all(v is None or v == "" for v in values):
                        continue
                    count += 1
                    raw += 1
                    source = {"file": file.name, "sheet": ws.title, "row": index}
                    row = {field: values[i] if i < len(values) else None for i, field in mapping.items()}
                    errors = []
                    for field in schema["required"]:
                        if row.get(field) is None or str(row[field]).strip() == "":
                            errors.append(f"missing:{field}")
                    for field in ("order_id", "product"):
                        row[field] = str(row.get(field) or "").strip()
                    if row["order_id"] and not re.fullmatch(r"[A-Za-z0-9_-]{1,80}", row["order_id"]):
                        errors.append("invalid:order_id")
                    try:
                        row["date"] = parse_date(row.get("date"))
                    except ValueError:
                        errors.append("invalid:date")
                    region = schema["regions"].get(str(row.get("region") or "").strip())
                    if region is None:
                        errors.append("invalid:region")
                    row["region"] = region
                    for field, maximum in (("quantity", schema["max_quantity"]), ("unit_price", schema["max_unit_price"])):
                        try:
                            value = number(row.get(field))
                            if value <= 0 or value > maximum or (field == "quantity" and value != value.to_integral_value()):
                                raise ValueError("outside allowed range")
                            if field == "unit_price" and value != value.quantize(CENT):
                                raise ValueError("price has more than two decimal places")
                            row[field] = int(value) if field == "quantity" else value.quantize(CENT)
                        except (ValueError, InvalidOperation):
                            errors.append(f"invalid:{field}")
                    if not any(e.endswith(":quantity") or e.endswith(":unit_price") for e in errors):
                        amount = (row["unit_price"] * row["quantity"]).quantize(CENT, rounding=ROUND_HALF_UP)
                        if row.get("amount") not in (None, ""):
                            try:
                                if abs(number(row["amount"]) - amount) > CENT:
                                    errors.append("mismatch:amount")
                            except ValueError:
                                errors.append("invalid:amount")
                        row["amount"] = amount
                    if errors:
                        quarantined += 1
                        issues.append({**source, "order_id": row["order_id"], "action": "quarantined", "reason": ";".join(errors)})
                    else:
                        candidates.append({**row, **source})
                sources.append({"file": file.name, "sheet": ws.title, "header_row": header, "raw_rows": count})
            if not matched:
                raise ValueError(f"No valid header within first 20 rows: {file.name}; required: {schema['required']}")
        finally:
            wb.close()
    groups = defaultdict(list)
    for row in candidates:
        groups[row["order_id"]].append(row)
    accepted = []
    for group in groups.values():
        signatures = {tuple(r[f] for f in FIELDS) for r in group}
        if len(signatures) > 1:
            quarantined += len(group)
            for row in group:
                issues.append({k: row[k] for k in ("file", "sheet", "row", "order_id")} | {"action": "quarantined", "reason": "conflicting:order_id"})
        else:
            accepted.append(group[0])
            for row in group[1:]:
                duplicates += 1
                issues.append({k: row[k] for k in ("file", "sheet", "row", "order_id")} | {"action": "duplicate_removed", "reason": "duplicate:order_id"})
    result = Result(accepted, issues, sources, raw, duplicates, quarantined)
    if not result.metrics["reconciled"]:
        raise ValueError("Row reconciliation failed")
    return result
