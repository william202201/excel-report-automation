"""Native Excel charts and an offline HTML preview from the same integer-cent totals."""
import csv
from datetime import date, datetime
from decimal import Decimal
from html import escape
import json
import os
from pathlib import Path
from tempfile import TemporaryDirectory

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .pipeline import FIELDS, Result


def safe_csv(value):
    text = str(value)
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text


def append(ws, values):
    # Explicit strings keep source text from becoming Excel formulas.
    ws.append([float(v) if isinstance(v, Decimal) else v for v in values])
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"
        elif isinstance(cell.value, (date, datetime)):
            cell.number_format = "yyyy-mm-dd"


def style(ws, table_name=None):
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor="18324B")
        c.font = Font(name="Calibri", color="FFFFFF", bold=True, size=11)
        c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    for column in ws.columns:
        width = max(len(str(c.value or "")) for c in list(column)[:300])
        ws.column_dimensions[column[0].column_letter].width = min(46, max(16, width + 3))
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
        if table_name:
            table = Table(displayName=table_name, ref=ws.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = '#,##0.00'


def pivots(result):
    frame = pd.DataFrame([{"month": r["date"].strftime("%Y-%m"), "region": r["region"],
                           "product": r["product"], "cents": int(r["amount"] * 100),
                           "quantity": r["quantity"]} for r in result.rows],
                         columns=["month", "region", "product", "cents", "quantity"])
    if frame.empty:
        return [], [], []
    monthly = frame.groupby("month", sort=True).agg(cents=("cents", "sum"), orders=("cents", "size"), quantity=("quantity", "sum"))
    regions = frame.groupby("region", sort=True)["cents"].sum()
    matrix = pd.pivot_table(frame, index="region", columns="product", values="cents", aggfunc="sum", fill_value=0, sort=True)
    return ([(i, int(r.cents), int(r.orders), int(r.quantity)) for i, r in monthly.iterrows()],
            [(i, int(v)) for i, v in regions.items()], matrix)


def render_html(result, monthly, regions):
    m = result.metrics
    max_month = max((v for _, v, _, _ in monthly), default=1) or 1
    bars = []
    for i, (label, cents, _, _) in enumerate(monthly):
        x, h = 30 + i * 57, cents / max_month * 160
        bars.append(f'<rect x="{x}" y="{200-h:.1f}" width="32" height="{h:.1f}" rx="4" fill="#0f9488"/><text x="{x+16}" y="225" text-anchor="middle">{label[-2:]}</text>')
    region_rows = "".join(f'<tr><td>{escape(label)}</td><td>¥{cents/100:,.2f}</td></tr>' for label, cents in regions)
    issues = "".join(f'<tr><td>{escape(str(r["file"]))}</td><td>{r["row"]}</td><td>{escape(r["reason"])}</td></tr>' for r in result.issues[:8])
    return f'''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Excel 报表自动化 · 运行结果</title><style>
*{{box-sizing:border-box}}body{{margin:0;background:#f1f5f8;color:#18324b;font:15px/1.65 -apple-system,BlinkMacSystemFont,"Segoe UI",sans-serif}}main{{max-width:1180px;margin:auto;padding:42px 32px}}.eyebrow{{font-size:12px;letter-spacing:2px;color:#547085}}h1{{margin:8px 0;font-size:34px;line-height:1.3}}p{{color:#637689}}.tag{{background:#d9f1e9;color:#126653;border-radius:20px;padding:5px 12px;font-size:12px}}header{{display:flex;justify-content:space-between;align-items:center}}.cards{{display:grid;grid-template-columns:repeat(4,1fr);gap:16px;margin:28px 0}}article,.panel{{background:white;border:1px solid #dce5ed;border-radius:12px;padding:22px}}article span{{color:#637689;font-size:13px}}article strong{{display:block;font-size:28px;margin-top:6px}}.grid{{display:grid;grid-template-columns:2fr 1fr;gap:18px}}h2{{font-size:17px;margin:0 0 16px}}svg{{width:100%;height:260px}}svg text{{font-size:11px;fill:#718396}}table{{border-collapse:collapse;width:100%;font-size:13px}}td,th{{text-align:left;padding:10px 8px;border-bottom:1px solid #edf1f4}}th{{color:#718396;font-weight:500}}.issues{{margin-top:18px}}.foot{{display:flex;justify-content:space-between;margin-top:22px;font-size:12px;color:#637689}}a{{color:#087f75}}@media(max-width:760px){{.cards{{grid-template-columns:1fr 1fr}}.grid{{grid-template-columns:1fr}}header{{display:block}}}}
</style><main><header><div><div class="eyebrow">EXCEL REPORT AUTOMATION / DEMO</div><h1>销售报表，自动汇总</h1><p>统一字段 · 异常隔离 · 订单去重 · 来源追溯</p></div><span class="tag">输入行数对账通过</span></header>
<section class="cards"><article><span>合并文件</span><strong>{m['input_files']} 份</strong></article><article><span>有效订单</span><strong>{m['accepted_rows']:,}</strong></article><article><span>销售额 / 人民币</span><strong>¥{Decimal(m['sales_total_cny']):,.2f}</strong></article><article><span>隔离异常 / 删除重复</span><strong>{m['quarantined_rows']} / {m['duplicate_rows']}</strong></article></section>
<section class="grid"><div class="panel"><h2>月度销售额</h2><svg viewBox="0 0 730 245" role="img" aria-label="按月汇总销售额的柱状图"><path d="M20 201H720" stroke="#dce5ed"/>{''.join(bars)}</svg></div><div class="panel"><h2>地区销售额</h2><table><thead><tr><th>地区</th><th>销售额</th></tr></thead><tbody>{region_rows}</tbody></table><p>仅汇总通过校验的订单。</p></div></section>
<section class="panel issues"><h2>问题清单预览 <small style="color:#718396;font-weight:400">前 8 条 / 共 {len(result.issues)} 条</small></h2><table><thead><tr><th>来源文件</th><th>Excel 行号</th><th>原因</th></tr></thead><tbody>{issues}</tbody></table></section>
<div class="foot"><span>模拟数据 · {m['raw_rows']:,} 输入行 = {m['accepted_rows']:,} 有效 + {m['quarantined_rows']} 隔离 + {m['duplicate_rows']} 重复</span><span><a href="report.xlsx">下载 Excel</a> · <a href="issues.csv">查看完整问题清单</a></span></div></main></html>'''


def write_report(result: Result, output: Path) -> dict:
    monthly, regions, matrix = pivots(result)
    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "总览"
    append(dashboard, ["指标", "结果"])
    for key, value in result.metrics.items():
        append(dashboard, [key, value])
    append(dashboard, ["数据说明", "模拟数据；金额单位 CNY；汇总表由 pandas 生成"])
    append(dashboard, ["刷新方式", "修改输入后重新运行命令；不是 Excel 原生交互透视表"])
    details = wb.create_sheet("清洗明细")
    append(details, FIELDS + ["file", "sheet", "row"])
    for r in result.rows:
        append(details, [r.get(k) for k in FIELDS + ["file", "sheet", "row"]])
    month_ws = wb.create_sheet("月度汇总")
    append(month_ws, ["月份", "销售额_CNY", "订单数", "销售数量"])
    for label, cents, count, quantity in monthly:
        append(month_ws, [label, cents / 100, count, quantity])
    region_ws = wb.create_sheet("地区汇总")
    append(region_ws, ["地区", "销售额_CNY"])
    for label, cents in regions:
        append(region_ws, [label, cents / 100])
    pivot_ws = wb.create_sheet("地区产品透视汇总")
    if hasattr(matrix, "columns"):
        append(pivot_ws, ["地区", *matrix.columns, "合计_CNY"])
        for region, row in matrix.iterrows():
            append(pivot_ws, [region, *(int(v)/100 for v in row), int(row.sum())/100])
        append(pivot_ws, ["合计", *(int(matrix[c].sum())/100 for c in matrix.columns), int(matrix.to_numpy().sum())/100])
    else:
        append(pivot_ws, ["地区", "合计_CNY"])
    issue_ws = wb.create_sheet("问题清单")
    issue_fields = ["file", "sheet", "row", "order_id", "action", "reason"]
    append(issue_ws, issue_fields)
    for issue in result.issues:
        append(issue_ws, [issue[k] for k in issue_fields])
    source_ws = wb.create_sheet("来源审计")
    append(source_ws, ["file", "sheet", "header_row", "raw_rows"])
    for source in result.sources:
        append(source_ws, list(source.values()))
    for i, ws in enumerate(wb):
        style(ws, f"Data{i}" if ws != dashboard else None)
    dashboard.column_dimensions["A"].width = 28
    dashboard.column_dimensions["B"].width = 70
    for cls, source, title, anchor in ((LineChart, month_ws, "月度销售额 / CNY", "D2"), (BarChart, region_ws, "地区销售额 / CNY", "D18")):
        if source.max_row > 1:
            chart = cls()
            chart.title = title
            chart.add_data(Reference(source, min_col=2, min_row=1, max_row=source.max_row), titles_from_data=True)
            chart.set_categories(Reference(source, min_col=1, min_row=2, max_row=source.max_row))
            chart.height, chart.width = 8, 20
            chart.legend = None
            dashboard.add_chart(chart, anchor)
    output.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(dir=output, prefix=".stage-") as temp:
        stage = Path(temp)
        wb.save(stage / "report.xlsx")
        with (stage / "issues.csv").open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=issue_fields)
            writer.writeheader()
            writer.writerows({k: safe_csv(v) for k, v in row.items()} for row in result.issues)
        (stage / "dashboard.html").write_text(render_html(result, monthly, regions), encoding="utf-8")
        (stage / "summary.json").write_text(json.dumps(result.metrics, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        for name in ("report.xlsx", "issues.csv", "dashboard.html", "summary.json"):
            os.replace(stage / name, output / name)
    return result.metrics
