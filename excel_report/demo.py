"""Deterministic synthetic fixtures: 12 deliberately different exports."""
from datetime import datetime, timedelta
from pathlib import Path
import random

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

PROFILES = [
    ["订单编号", "订单日期", "地区", "产品", "数量", "单价", "销售额"],
    ["Order ID", "Date", "Region", "Product", "Qty", "Unit Price", "Amount"],
    ["订单号", "日期", "区域", "商品名称", "销售数量", "销售单价", "金额"],
    ["order_id", "order_date", "sales_region", "item", "quantity", "unit_price", "total_amount"],
]


def generate(folder: Path, rows_per_file: int = 500) -> None:
    if rows_per_file < 60:
        raise ValueError("rows_per_file must be at least 60")
    folder.mkdir(parents=True, exist_ok=True)
    if any(folder.iterdir()):
        raise ValueError(f"Refusing to overwrite a non-empty fixture directory: {folder}")
    rng = random.Random(20261018)
    for month in range(1, 13):
        wb = Workbook()
        cover = wb.active
        cover.title = "说明"
        cover.append(["SYNTHETIC DATA — fictional sales for a portfolio demo"])
        ws = wb.create_sheet("Sales" if month % 2 else "销售明细")
        for _ in range(month % 3):
            ws.append(["模拟分支机构销售导出；单位：人民币元"])
        profile = PROFILES[(month - 1) % 4]
        order = list(range(7))
        rng.shuffle(order)
        ws.append([profile[i] for i in order])
        for c in ws[ws.max_row]:
            c.font = Font(color="FFFFFF", bold=True)
            c.fill = PatternFill("solid", fgColor="18324B")
        for i in range(rows_per_file):
            qty, price = rng.randint(1, 12), rng.choice([19.9, 49.5, 99, 129, 299])
            day = datetime(2025, month, 1) + timedelta(days=i % 27)
            date = [day, day.strftime("%Y/%m/%d"), day.strftime("%Y年%m月%d日")][month % 3]
            region = rng.choice(["华东", "华南", "华北", "西部"])
            if month % 2:
                region = {"华东": "East", "华南": "South", "华北": "North", "西部": "West"}[region]
            row = [f"SO-{month:02d}-{i:05d}", date, region,
                   rng.choice(["键盘", "鼠标", "显示器", "扩展坞"]), qty, price, round(qty * price, 2)]
            if month % 4 == 0:
                row[5] = f"￥{price:,.2f}"
            if i == 5:
                row[1] = "2025-02-30"
            elif i == 11:
                row[4] = -3
            elif i == 17:
                row[3] = None
            elif i == 23:
                row[5] = "not-a-number"
            elif i == 29:
                row[2] = "未知地区"
            elif i == 35:
                row[6] = 0.01
            ws.append([row[j] for j in order])
            if i == 41:
                ws.append([row[j] for j in order])  # exact duplicate
            if i == 47:
                conflict = row.copy()
                conflict[3] = "不同产品"
                ws.append([conflict[j] for j in order])  # same ID, different content
        ws.freeze_panes = f"A{month % 3 + 2}"
        for col in "ABCDEFG":
            ws.column_dimensions[col].width = 22
        wb.save(folder / f"branch_{month:02d}.xlsx")
